"""
The free templates behind an e-mail address (D-241): Excel and Word files built from what the site
already says, so a template and the page it comes from cannot disagree.

  python3 scripts/maler/build.py        writes private/maler/*

Reads messages/no.json (the law table on /lovkrav, the risk assessment and «hvor ofte» articles, the
playbook and the factors) and scripts/maler/qpsnordic34.json (QPSNordic 34+, verbatim, X-107). The
files are Norwegian, as the law and the questionnaire are. Needs openpyxl and python-docx; the outputs
are committed, like the social cards, so neither the build nor CI needs Python. Run it again when one
of those texts changes.
"""
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'private' / 'maler'
NO = json.loads((ROOT / 'messages' / 'no.json').read_text())
QPS = json.loads((ROOT / 'scripts' / 'maler' / 'qpsnordic34.json').read_text())

# the site's tokens (tailwind.config.ts)
INK, BODY, MUT = '191510', '3A342A', '5F5849'
AC, SBG, MINT, LINE, SF, PEACH = 'F5C64A', 'FBEBBE', 'CFE7E4', 'E8DFC9', 'FFFDF6', 'FBD5C4'
FOOT = 'Laget av Orgpuls · orgpuls.no · Fri til bruk i egen virksomhet. Malen er et hjelpemiddel, ikke juridisk rådgivning.'


def plain(s: str) -> str:
    """A message as plain text: links to their words, word joiners and no-break spaces out."""
    s = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', s)
    s = s.replace('⁠', '').replace(' ', ' ').replace('**', '')
    return s


def block(article: str, t: str, after: str | None = None):
    """The first block of type t in an article (after the h2 with text `after`, if given)."""
    blocks = NO['seo']['articles'][article]['blocks']
    start = 0
    if after:
        start = next(i for i, b in enumerate(blocks) if b['t'] == 'h2' and b['text'] == after)
    return next(b for b in blocks[start:] if b['t'] == t)


# --------------------------------------------------------------------------- excel helpers
thin = Side(style='thin', color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')


def title(ws, text: str, sub: str | None = None, width: int = 8):
    ws['A1'] = text
    ws['A1'].font = Font(name='Calibri', size=18, bold=True, color=INK)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.row_dimensions[1].height = 30
    if sub:
        ws['A2'] = sub
        ws['A2'].font = Font(name='Calibri', size=11, color=MUT)
        ws['A2'].alignment = WRAP
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        ws.row_dimensions[2].height = 32
    ws.sheet_view.showGridLines = False


def header_row(ws, row: int, labels: list[str], fill: str = INK, color: str = 'FFFFFF'):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(name='Calibri', bold=True, color=color)
        c.fill = PatternFill('solid', fgColor=fill)
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = BOX
    ws.row_dimensions[row].height = 32


def cell(ws, row: int, col: int, value, bold: bool = False, fill: str | None = None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name='Calibri', bold=bold, color=BODY if not bold else INK)
    c.alignment = WRAP
    c.border = BOX
    if fill:
        c.fill = PatternFill('solid', fgColor=fill)
    return c


def widths(ws, ws_widths: list[int]):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def text_sheet(ws, heading: str, paras: list[tuple[str, str]]):
    """A sheet of prose: (kind, text) with kind h, p, li or src."""
    title(ws, heading, width=1)
    ws.column_dimensions['A'].width = 86
    r = 3
    for kind, text in paras:
        c = ws.cell(row=r, column=1, value=('•  ' + text) if kind == 'li' else text)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if kind == 'h':
            c.font = Font(name='Calibri', size=13, bold=True, color=INK)
            r += 1
        elif kind == 'src':
            c.font = Font(name='Calibri', size=9, italic=True, color=MUT)
        else:
            c.font = Font(name='Calibri', size=11, color=BODY)
        lines = max(1, len(text) // 82 + 1)
        ws.row_dimensions[r].height = 15 * lines + 4
        r += 1
    ws.cell(row=r + 1, column=1, value=FOOT).font = Font(name='Calibri', size=9, italic=True, color=MUT)


def foot(ws, row: int):
    ws.cell(row=row, column=1, value=FOOT).font = Font(name='Calibri', size=9, italic=True, color=MUT)


# --------------------------------------------------------------------------- 1 · risikovurdering
def risikovurdering(path: Path):
    wb = Workbook()
    law = next(b for b in NO['seo']['lp']['lovkrav']['blocks'] if b['t'] == 'table')
    quote = block('risikovurdering', 'quote')
    steps = block('risikovurdering', 'timeline', 'En arbeidsmåte i fire steg')['items']
    three = block('risikovurdering', 'cards', 'Hvilke tre forhold skal vurderes?')['items']
    document = block('risikovurdering', 'ul', 'Hva skal dokumenteres?')['items']

    ws = wb.active
    ws.title = 'Slik bruker du malen'
    paras = [
        ('p', 'Malen hjelper dere å gjøre og dokumentere risikovurderingen av det psykososiale arbeidsmiljøet, slik '
              'arbeidsmiljøloven § 3-1 og forskrift om utførelse av arbeid kapittel 1A krever. Den bygger på en kartlegging '
              'dere allerede har gjort – en medarbeiderundersøkelse, vernerunder, samtaler eller avviksmeldinger.'),
        ('h', 'Hva forskriften sier'),
        ('p', f'«{plain(quote["text"])}» ({plain(quote["cite"])})'),
        ('p', 'Vurderingen gjøres i samarbeid med de ansatte og deres representanter, tilpasses virksomhetens art og størrelse, '
              'gjentas regelmessig og ved endringer, og dokumenteres (forskriften § 1A-2).'),
        ('h', 'Tre forhold som skal vurderes'),
        *[('li', f'{plain(c["title"])}: {plain(c["text"])}') for c in three],
        ('h', 'Arbeidsmåte i fire steg'),
        *[('li', f'{s["date"]} – {s["title"]}{" " if s.get("text", "")[:1].islower() else ": "}{plain(s.get("text", ""))}') for s in steps],
        ('h', 'Slik fyller du ut arket «Risikovurdering»'),
        ('li', 'Én rad per faktor. Faktorene og hjemlene er de arbeidsmiljøloven og forskriften nevner.'),
        ('li', 'Skriv hva kartleggingen viser, hvor mange det gjelder, og hvor lenge det har vart.'),
        ('li', 'Sett sannsynlighet og konsekvens fra 1 til 4 (se arket «Risikomatrise»). Risikoen regnes ut, og nivået fargelegges.'),
        ('li', 'Noter hvilke andre faktorer som forsterker eller demper risikoen, og begrunn konklusjonen.'),
        ('li', 'Faktorer med middels eller høy risiko får et tiltak med ansvarlig og frist. Før tiltakene videre i arket «Tiltaksplan».'),
        ('h', 'Dette bør dokumentasjonen vise'),
        *[('li', plain(d)) for d in document],
        ('h', 'Skalaene er et forslag'),
        ('p', 'Forskriften angir ingen terskelverdier eller faste skalaer. Skalaene for sannsynlighet og konsekvens i denne malen '
              'er et forslag til en felles måte å vurdere på; tilpass dem til virksomheten, og bruk de samme hver gang.'),
        ('src', 'Kilder: arbeidsmiljøloven § 3-1 og § 4-3, forskrift om utførelse av arbeid kapittel 1A (Lovdata), og Arbeidstilsynets '
                'kommentarer til forskriften. Se også orgpuls.no/artikler/risikovurdering-av-psykososialt-arbeidsmiljo.'),
    ]
    text_sheet(ws, 'Risikovurdering av det psykososiale arbeidsmiljøet', paras)

    ws = wb.create_sheet('Risikovurdering')
    title(ws, 'Risikovurdering', 'Fyll inn én rad per faktor. Sannsynlighet og konsekvens settes fra 1 til 4; risiko og nivå regnes ut.', 14)
    meta = ['Virksomhet', 'Enhet eller gruppe', 'Dato for vurderingen', 'Deltakere (leder, verneombud, ansatte)', 'Grunnlag (kartlegging og dato)']
    for i, m in enumerate(meta):
        cell(ws, 3 + i, 1, m, bold=True, fill=SBG)
        ws.merge_cells(start_row=3 + i, start_column=2, end_row=3 + i, end_column=6)
        cell(ws, 3 + i, 2, '')
    head = ['Faktor', 'Hjemmel', 'Hva viser kartleggingen?', 'Hvor mange gjelder det?', 'Hvor lenge har det vart?',
            'Sannsynlighet (1–4)', 'Konsekvens (1–4)', 'Risiko (S × K)', 'Risikonivå', 'Forsterkes eller dempes av',
            'Begrunnelse for konklusjonen', 'Tiltak', 'Ansvarlig', 'Frist']
    top = 9
    header_row(ws, top, head)
    dv = DataValidation(type='whole', operator='between', formula1='1', formula2='4', allow_blank=True,
                        error='Skriv et tall fra 1 til 4.', errorTitle='Fra 1 til 4')
    ws.add_data_validation(dv)
    rows = law['rows']
    for i, r in enumerate(rows):
        row = top + 1 + i
        cell(ws, row, 1, plain(r[0]), bold=True)
        cell(ws, row, 2, plain(r[1]))
        for col in range(3, 15):
            cell(ws, row, col, None)
        ws.cell(row=row, column=8, value=f'=IF(AND(ISNUMBER(F{row}),ISNUMBER(G{row})),F{row}*G{row},"")')
        ws.cell(row=row, column=9, value=f'=IF(H{row}="","",IF(H{row}>=9,"Høy",IF(H{row}>=4,"Middels","Lav")))')
        dv.add(f'F{row}:G{row}')
        ws.row_dimensions[row].height = 48
    last = top + len(rows)
    # the overall risk: § 1A-2 third paragraph
    row = last + 1
    cell(ws, row, 1, 'Samlet vurdering', bold=True, fill=SBG)
    cell(ws, row, 2, 'forskriften § 1A-2 (3)', fill=SBG)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=14)
    cell(ws, row, 3, 'Hvordan påvirker faktorene hverandre, og utgjør de samlet en risiko? Flere middels svake faktorer i samme gruppe kan til sammen gi høyere risiko.')
    ws.row_dimensions[row].height = 48
    rng = f'I{top + 1}:I{last}'
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Høy"'], fill=PatternFill('solid', fgColor='F0B9A0')))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Middels"'], fill=PatternFill('solid', fgColor=AC)))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"Lav"'], fill=PatternFill('solid', fgColor=MINT)))
    widths(ws, [30, 20, 34, 16, 16, 12, 12, 11, 11, 26, 34, 34, 16, 12])
    ws.freeze_panes = ws.cell(row=top + 1, column=2)
    foot(ws, row + 2)

    ws = wb.create_sheet('Risikomatrise')
    title(ws, 'Risikomatrise', 'Risiko = sannsynlighet × konsekvens. 1–3 lav, 4–8 middels, 9–16 høy. Skalaene er et forslag; bruk de samme hver gang.', 6)
    cell(ws, 4, 1, 'Sannsynlighet ↓ / Konsekvens →', bold=True, fill=SBG)
    for k in range(1, 5):
        cell(ws, 4, 1 + k, k, bold=True, fill=SBG)
    for s in range(4, 0, -1):
        r = 9 - s
        cell(ws, r, 1, s, bold=True, fill=SBG)
        for k in range(1, 5):
            v = s * k
            cell(ws, r, 1 + k, v, fill='F0B9A0' if v >= 9 else AC if v >= 4 else MINT).alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[r].height = 28
    scales = [
        ('Sannsynlighet', ['1 – Lite sannsynlig: få kjenner seg igjen, og det er ingen tegn til at det øker',
                           '2 – Mulig: noen kjenner seg igjen, eller det har skjedd før',
                           '3 – Sannsynlig: mange kjenner seg igjen, eller det har vart over tid',
                           '4 – Svært sannsynlig: de fleste kjenner seg igjen, eller det skjer nå']),
        ('Konsekvens', ['1 – Liten: ubehag som går over av seg selv',
                        '2 – Moderat: belastning som påvirker trivsel og innsats over tid',
                        '3 – Alvorlig: fare for helseplager, sykefravær eller at folk slutter',
                        '4 – Svært alvorlig: fare for varig helseskade, vold, trakassering eller annen alvorlig skade']),
    ]
    r = 11
    for name, lines in scales:
        cell(ws, r, 1, name, bold=True, fill=SBG)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        r += 1
        for line in lines:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            cell(ws, r, 1, line)
            r += 1
        r += 1
    widths(ws, [34, 14, 14, 14, 14, 14])
    foot(ws, r + 1)

    ws = wb.create_sheet('Tiltaksplan')
    title(ws, 'Tiltaksplan', 'Tiltakene risikovurderingen førte til, med ansvarlig, frist og hvordan dere ser om de virket.', 9)
    head = ['Tiltak', 'Faktor', 'Hva skal bli bedre?', 'Ansvarlig', 'Frist', 'Status', 'Følges opp med', 'Evaluert (dato)', 'Virket det?']
    header_row(ws, 4, head)
    sv = DataValidation(type='list', formula1='"Ikke startet,Pågår,Fullført,Avbrutt"', allow_blank=True)
    ws.add_data_validation(sv)
    for row in range(5, 25):
        for col in range(1, 10):
            cell(ws, row, col, None)
        sv.add(f'F{row}')
        ws.row_dimensions[row].height = 30
    widths(ws, [34, 24, 30, 16, 12, 14, 30, 14, 24])
    foot(ws, 26)
    wb.active = 1
    wb.save(path)


# --------------------------------------------------------------------------- 2 · tiltaksbank
def tiltaksbank(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Slik bruker du malen'
    text_sheet(ws, 'Tiltaksbank for det psykososiale arbeidsmiljøet', [
        ('p', 'Tre konkrete tiltak for hver av de elleve faktorene i Orgpuls\' spørsmålssett – 33 i alt – med hvor lang tid de tar, '
              'hvordan dere gjør dem, og forskningen de bygger på.'),
        ('h', 'Slik bruker du den'),
        ('li', 'Start med faktorene som står svakest i kartleggingen eller risikovurderingen.'),
        ('li', 'Velg to eller tre tiltak, ikke ti. Gi hvert tiltak en ansvarlig og en frist.'),
        ('li', 'Kolonnen «Følg med på» er påstanden i spørsmålssettet som bør bevege seg hvis tiltaket virker. Mål den igjen etter noen måneder.'),
        ('li', 'Bestem tiltakene sammen med de ansatte. Tiltak virker best når de som berøres, er med på å finne dem.'),
        ('h', 'Typer tiltak'),
        ('li', 'Workshop: en samling der teamet finner løsninger sammen.'),
        ('li', 'Rutine: en fast praksis som gjentas.'),
        ('li', 'Lederpraksis: noe lederen gjør annerledes.'),
        ('src', 'Kilde: Orgpuls\' tiltaksbibliotek. Se også orgpuls.no/artikler/fra-resultater-til-tiltak.'),
    ])
    ws = wb.create_sheet('Tiltak')
    title(ws, 'Tiltaksbank', '33 tiltak, tre per faktor. Filtrer på faktor eller type.', 7)
    head = ['Faktor', 'Tiltak', 'Type', 'Tid', 'Slik gjør dere det', 'Følg med på (påstand)', 'Grunnlag']
    header_row(ws, 4, head)
    reg = (ROOT / 'lib' / 'playbook' / 'registry.ts').read_text()
    rows = re.findall(r"\['(\w+)', ([123]), '(\w+)', (\d)\]", reg)
    r = 5
    for factor, n, kind, watch in rows:
        p = NO['playbook'][factor]
        m = p[f'm{n}']
        f = NO['factor'][factor]
        cell(ws, r, 1, f['label'], bold=True)
        cell(ws, r, 2, plain(m['title']), bold=True)
        cell(ws, r, 3, NO['playbook']['kind'][kind])
        cell(ws, r, 4, plain(m['time']))
        cell(ws, r, 5, plain(m['how']))
        cell(ws, r, 6, f'«{plain(f["s" + watch])}»')
        cell(ws, r, 7, plain(p['ev']))
        ws.row_dimensions[r].height = 76
        r += 1
    ws.auto_filter.ref = f'A4:G{r - 1}'
    ws.freeze_panes = 'C5'
    widths(ws, [24, 30, 13, 14, 50, 34, 50])
    foot(ws, r + 1)
    wb.active = 1
    wb.save(path)


# --------------------------------------------------------------------------- 3 · årshjul
def arshjul(path: Path):
    wb = Workbook()
    wheel = block('hvorOfte', 'timeline', 'Et eksempel på et årshjul')['items']
    changes = block('hvorOfte', 'ul', 'Når utløser en endring ny kartlegging?')['items']
    tips = block('hvorOfte', 'ul', 'Hvordan holde svarviljen oppe')['items']
    ws = wb.active
    ws.title = 'Slik bruker du malen'
    text_sheet(ws, 'Årshjul for arbeidet med det psykososiale arbeidsmiljøet', [
        ('p', 'Forskriften sier at kartleggingen og risikovurderingen skal gjentas regelmessig og ved endringer som kan ha betydning '
              'for risikoen (forskrift om utførelse av arbeid § 1A-2). Den setter ikke et fast intervall. Årshjulet er et eksempel '
              'på en rytme: én hovedkartlegging i året, og korte oppfølgingsmålinger der tiltakene løper.'),
        ('h', 'Slik bruker du det'),
        ('li', 'Arket «Årshjul» har eksempelet ferdig utfylt. Flytt månedene så de passer virksomheten, og unngå feriene.'),
        ('li', 'Fyll inn hvem som har ansvaret, og hvem som skal være med – verneombudet og de ansatte skal være med på kartleggingen og risikovurderingen.'),
        ('li', 'Arket «Ved endringer» lister endringer som bør utløse en ny kartlegging utenom årshjulet.'),
        ('h', 'Slik holder dere svarviljen oppe'),
        *[('li', plain(t)) for t in tips],
        ('src', 'Kilder: forskrift om utførelse av arbeid § 1A-2 (Lovdata). Se også orgpuls.no/artikler/hvor-ofte-bor-dere-male-arbeidsmiljoet.'),
    ])
    ws = wb.create_sheet('Årshjul')
    title(ws, 'Årshjul', 'Eksempelet er fylt inn. Tilpass måneder, ansvar og deltakere.', 6)
    header_row(ws, 4, ['Måned', 'Aktivitet', 'Hvem har ansvaret', 'Hvem er med', 'Hjemmel', 'Gjort (dato)'])
    months = ['Januar', 'Februar', 'Mars', 'April', 'Mai', 'Juni', 'Juli', 'August', 'September', 'Oktober', 'November', 'Desember']
    plan = {w['date']: plain(w['title']) for w in wheel}
    legal = {'Februar': 'aml § 3-1 (2) c · forskriften § 1A-2', 'Mars': 'forskriften § 1A-2 · aml § 3-1 (2) c'}
    for i, mth in enumerate(months):
        r = 5 + i
        has = mth in plan
        cell(ws, r, 1, mth, bold=True, fill=SBG if has else None)
        cell(ws, r, 2, plan.get(mth, ''), fill=SBG if has else None)
        cell(ws, r, 3, '')
        cell(ws, r, 4, 'Leder, verneombud og de ansatte' if has else '')
        cell(ws, r, 5, legal.get(mth, ''))
        cell(ws, r, 6, '')
        ws.row_dimensions[r].height = 34
    widths(ws, [12, 56, 22, 28, 30, 14])
    foot(ws, 19)
    ws = wb.create_sheet('Ved endringer')
    title(ws, 'Endringer som bør utløse en ny kartlegging', 'Forskriften: kartleggingen skal gjentas ved endringer som kan ha betydning for risikoforholdene.', 4)
    header_row(ws, 4, ['Endring', 'Gjelder det oss nå?', 'Når kartlegger vi?', 'Ansvarlig'])
    for i, c in enumerate(changes):
        r = 5 + i
        cell(ws, r, 1, plain(c)[0].upper() + plain(c)[1:], bold=True)
        for col in (2, 3, 4):
            cell(ws, r, col, '')
        ws.row_dimensions[r].height = 30
    widths(ws, [56, 20, 20, 20])
    foot(ws, 5 + len(changes) + 1)
    wb.active = 1
    wb.save(path)


# --------------------------------------------------------------------------- 4 · QPSNordic 34+
def shade(c, hex_fill: str):
    tcPr = c._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def doc_base(heading: str) -> Document:
    d = Document()
    st = d.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor.from_string(BODY)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
        fp = s.footer.paragraphs[0]
        fp.text = FOOT
        fp.runs[0].font.size = Pt(8)
        fp.runs[0].font.color.rgb = RGBColor.from_string(MUT)
    h = d.add_heading(heading, level=0)
    for r in h.runs:
        r.font.color.rgb = RGBColor.from_string(INK)
    return d


def h(d: Document, text: str, level: int = 1):
    p = d.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = RGBColor.from_string(INK)


def small(d: Document, text: str):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    r.italic = True
    r.font.color.rgb = RGBColor.from_string(MUT)


def qps_docx(path: Path):
    d = doc_base('QPSNordic 34+')
    p = d.add_paragraph('Generelt nordisk spørreskjema for psykologiske og sosiale faktorer i arbeid – kortversjonen')
    p.runs[0].bold = True
    h(d, 'Til den som gjennomfører undersøkelsen', 2)
    for t in [
        'Spørsmålene er gjengitt ordrett fra brukerveiledningen til QPSNordic, utgitt av Statens arbeidsmiljøinstitutt (STAMI). '
        'Veiledningen sier at spørreskjemaet kan kopieres, og at man kan utvikle egne systemer for registrering, analyse og rapportering.',
        'Kortversjonen har ingen skalaer. Veiledningen understreker at testingen av validitet og reliabilitet som er gjort for '
        'QPSNordic, ikke gjelder for kortversjonen. Bruk resultatene per spørsmål.',
        'Veiledningen anbefaler ikke statistiske analyser for grupper med færre enn 10 personer. For 10–20 personer anbefales '
        'middelverdier og medianer, og svarene kan slås sammen til tre kategorier (1–2, 3 og 4–5). Fulle prosentfordelinger '
        'anbefales for grupper over 20.',
        'De ansatte skal på forhånd få vite hva undersøkelsen skal brukes til, hvordan den gjennomføres, og hvordan anonymiteten '
        'sikres. Det skal være frivillig å delta, og ingen resultater om enkeltpersoner skal presenteres.',
        'Svarregistreringen i Excel (egen fil) regner ut antall svar, middelverdi og fordeling per spørsmål.',
    ]:
        d.add_paragraph(t, style='List Bullet')
    small(d, 'Kilde: ' + QPS['source'])
    d.add_page_break()

    h(d, 'Besvarelse av spørreskjemaet', 1)
    for t in QPS['intro']:
        d.add_paragraph(t)
    h(d, 'Bakgrunnsspørsmål', 2)
    for b in QPS['background']:
        line = f'{b["id"]}. {b["text"]}'
        if b['options']:
            line += '   ' + '   '.join(f'☐ {o}' for o in b['options'])
        else:
            line += '   ______________________________'
        d.add_paragraph(line)

    current = None
    table = None
    for it in QPS['items']:
        key = (it['scale'], it.get('lead') if it.get('lead', '').startswith('Hvordan') else None)
        if key != current:
            current = key
            if it.get('lead'):
                lp = d.add_paragraph()
                lr = lp.add_run(it['lead'])
                lr.italic = True
                lr.bold = it['n'] == 37
            scale = QPS['scales'][it['scale']]
            table = d.add_table(rows=1, cols=7)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            hdr[0].text = 'Nr.'
            hdr[1].text = ''
            for i, s in enumerate(scale):
                hdr[2 + i].text = s
            for c in hdr:
                shade(c, SBG)
                for par in c.paragraphs:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in par.runs:
                        r.font.size = Pt(8)
                        r.bold = True
        row = table.add_row().cells
        row[0].text = str(it['n'])
        row[1].text = it['text']
        for i in range(5):
            row[2 + i].text = str(i + 1)
            row[2 + i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for c in row:
            for par in c.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(9.5)
        widths_cm = [1.0, 8.2, 1.6, 1.6, 1.6, 1.6, 1.6]
        for i, w in enumerate(widths_cm):
            table.columns[i].width = Cm(w)
        for r_ in table.rows:
            for i, w in enumerate(widths_cm):
                r_.cells[i].width = Cm(w)
        if it['n'] == 37:
            d.add_paragraph()
    d.add_paragraph()
    small(d, 'Kilde: ' + QPS['source'])
    d.save(path)


def qps_xlsx(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Slik bruker du malen'
    text_sheet(ws, 'QPSNordic 34+ – svarregistrering', [
        ('p', 'Registrer svarene fra papirskjemaene på arket «Svar»: én rad per skjema, ett tall fra 1 til 5 per spørsmål. '
              'La cellen stå tom hvis spørsmålet ikke er besvart. Arket «Oppsummering» regner ut resten.'),
        ('h', 'Før dere ser på tallene'),
        ('li', 'Veiledningen anbefaler ikke statistiske analyser for grupper under 10 personer. Oppsummeringen viser derfor ingenting før minst 10 har svart.'),
        ('li', 'For 10–20 personer anbefales middelverdi og median, og svarene slått sammen til tre kategorier: 1–2, 3 og 4–5. Det er det oppsummeringen viser.'),
        ('li', 'Kortversjonen har ingen skalaer, og testingen av validitet og reliabilitet gjelder ikke for den. Se på spørsmålene ett og ett.'),
        ('li', 'Ikke registrer navn eller noe annet som kan knytte et skjema til en person. Oppbevar papirskjemaene innelåst, og makuler dem når svarene er registrert.'),
        ('li', 'Høy verdi betyr «ofte» eller «mye» – det er bra på noen spørsmål (for eksempel nr. 17, støtte fra kollegaer) og dårlig på andre (for eksempel nr. 2, for mye å gjøre).'),
        ('src', 'Kilde: ' + QPS['source']),
    ])
    items = QPS['items']
    ws = wb.create_sheet('Svar')
    title(ws, 'Svar', 'Én rad per skjema. Tall fra 1 til 5; tomt hvis ubesvart.', 10)
    head = ['Skjema'] + [f'{it["n"]}. {it["text"]}' for it in items]
    header_row(ws, 4, head, fill=SBG, color=INK)
    ws.row_dimensions[4].height = 120
    for c in range(2, len(head) + 1):
        ws.cell(row=4, column=c).alignment = Alignment(wrap_text=True, vertical='bottom', text_rotation=90)
        ws.column_dimensions[get_column_letter(c)].width = 5.5
    ws.column_dimensions['A'].width = 9
    dv = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True,
                        error='Skriv et tall fra 1 til 5, eller la cellen stå tom.', errorTitle='Fra 1 til 5')
    ws.add_data_validation(dv)
    last = 4 + 300
    for r in range(5, last + 1):
        ws.cell(row=r, column=1, value=r - 4).font = Font(color=MUT)
    dv.add(f'B5:{get_column_letter(len(head))}{last}')
    ws.freeze_panes = 'B5'

    ws2 = wb.create_sheet('Oppsummering')
    title(ws2, 'Oppsummering', 'Vises først når minst 10 har svart på spørsmålet (veiledningen s. 14).', 8)
    header_row(ws2, 4, ['Nr.', 'Spørsmål', 'Antall svar', 'Middelverdi', 'Median', 'Andel 1–2', 'Andel 3', 'Andel 4–5'])
    for i, it in enumerate(items):
        r = 5 + i
        col = get_column_letter(2 + i)
        rng = f'Svar!{col}$5:{col}${last}'
        cell(ws2, r, 1, it['n'], bold=True)
        cell(ws2, r, 2, (it.get('lead', '') + ' ' if it.get('lead', '').startswith('Hvordan') else '') + it['text'])
        cell(ws2, r, 3, f'=COUNT({rng})')
        cell(ws2, r, 4, f'=IF(C{r}<10,"",ROUND(AVERAGE({rng}),2))')
        cell(ws2, r, 5, f'=IF(C{r}<10,"",MEDIAN({rng}))')
        cell(ws2, r, 6, f'=IF(C{r}<10,"",(COUNTIF({rng},1)+COUNTIF({rng},2))/C{r})')
        cell(ws2, r, 7, f'=IF(C{r}<10,"",COUNTIF({rng},3)/C{r})')
        cell(ws2, r, 8, f'=IF(C{r}<10,"",(COUNTIF({rng},4)+COUNTIF({rng},5))/C{r})')
        for c in (6, 7, 8):
            ws2.cell(row=r, column=c).number_format = '0%'
        ws2.row_dimensions[r].height = 30
    end = 4 + len(items)
    ws2.conditional_formatting.add(f'C5:C{end}', FormulaRule(formula=['AND(C5>0,C5<10)'], fill=PatternFill('solid', fgColor=PEACH)))
    widths(ws2, [6, 70, 12, 12, 10, 11, 10, 11])
    ws2.freeze_panes = 'C5'
    foot(ws2, end + 2)
    wb.active = 1
    wb.save(path)


# --------------------------------------------------------------------------- 5 · informasjonsskriv og sjekkliste
def informasjonsskriv(path: Path):
    d = doc_base('Informasjonsskriv og sjekkliste')
    d.add_paragraph('To maler for kartleggingen av det psykososiale arbeidsmiljøet: et skriv til de ansatte før undersøkelsen, '
                    'og en sjekkliste for dem som planlegger den. Tekst i [hakeparenteser] fylles inn eller fjernes.')
    h(d, 'Del 1 · Informasjonsskriv til de ansatte', 1)
    small(d, 'QPSNordic-veiledningen sier at de ansatte på forhånd skal få vite formålet, hvordan undersøkelsen gjennomføres, '
             'hva den spør om, hvordan resultatene analyseres og presenteres, og hvordan anonymiteten sikres – skriftlig og muntlig.')
    sections = [
        ('Hvorfor vi gjør dette',
         'Vi kartlegger det psykososiale arbeidsmiljøet hos [virksomhet] for å finne ut hva som fungerer, og hva som bør bli bedre. '
         'Arbeidsmiljøloven og forskrift om utførelse av arbeid krever at vi kartlegger, vurderer risikoen og setter inn tiltak – '
         'i samarbeid med dere og verneombudet. Svarene dine er grunnlaget for det arbeidet.'),
        ('Hvordan det foregår',
         'Undersøkelsen er åpen fra [dato] til [dato]. Den tar rundt [antall] minutter. Du svarer [på papir / på mobil eller PC via en lenke]. '
         'Spørsmålene handler om arbeidsmengde, roller, støtte, samarbeid, ledelse og hvordan du har det på jobb.'),
        ('Anonymitet',
         'Undersøkelsen er anonym. [Beskriv hvordan: for eksempel at skjemaet ikke har navn, at svarene registreres samlet, og at '
         'papirskjemaene makuleres etterpå.] Resultatene vises bare for grupper med minst [antall] svar, slik at ingen kan se '
         'hva den enkelte har svart. Ingen resultater om enkeltpersoner blir presentert.'),
        ('Frivillig',
         'Det er frivillig å svare, og du kan hoppe over spørsmål du ikke vil svare på. Vi håper likevel at så mange som mulig deltar, '
         'fordi resultatene blir mer pålitelige jo flere som svarer.'),
        ('Hva som skjer med svarene',
         'Svarene analyseres av [hvem]. Vi går gjennom resultatene sammen med dere på [møte/dato], og vurderer sammen med '
         'verneombudet hva som er viktigst å gjøre noe med. Tiltakene får en ansvarlig og en frist, og vi følger opp om de virker. '
         'Svarene oppbevares [hvor og hvor lenge] og slettes [når].'),
        ('Spørsmål',
         'Har du spørsmål om undersøkelsen, kan du ta kontakt med [navn og rolle] eller verneombudet [navn].'),
    ]
    for head, body in sections:
        h(d, head, 2)
        d.add_paragraph(body)
    d.add_paragraph('Vennlig hilsen\n[navn]\n[rolle]')
    d.add_page_break()

    h(d, 'Del 2 · Sjekkliste før kartleggingen', 1)
    h(d, 'Forhold som må vektlegges', 2)
    for t in [
        'Den øverste ledelsens forpliktelse om aktivt å støtte gjennomføring av undersøkelsen.',
        'Linjeledelsens forpliktelse til å bistå i å utføre de nødvendige forbedringene.',
        'De ansattes forpliktelse om å svare på spørreskjemaene, tolke resultatene og legge planer for forbedring.',
        'At stabsorganers rolle avklares.',
        'At fagforeningenes rolle avklares.',
        'At alle på avdelingsnivå deltar i undersøkelsen og i diskusjonene under tilbakemeldingen.',
        'At eventuelle andre utviklingstiltak eller endringer som er i gang i organisasjonen blir koordinert med QPSNordic undersøkelsen.',
        'Sikring av at etiske retningslinjer blir fulgt.',
        'Om undersøkelser bør bli utført flere ganger, og med hvilke intervaller.',
    ]:
        d.add_paragraph('☐  ' + t)
    h(d, 'Spørsmål som må avklares', 2)
    for t in [
        'Hva er formålet for undersøkelsen?',
        'Er deltakerne klar over, og godtar, målsetninger for undersøkelsen?',
        'Bør undersøkelsen gjennomføres samtidig i hele organisasjonen, eller vil det være bedre å ta for seg enhet for enhet?',
        'Vil det være behov for tilleggsskalaer for å måle helse, trivsel og lignende?',
        'Hvem tar ansvar for aktiviteten som helhet, og hva er hans/hennes oppgaver og myndighet?',
        'Hvordan vil analyser av data bli utført?',
        'Hvordan skal resultater presenteres i skriftlige rapporter?',
        'Når og hvor vil tilbakemeldinger bli gitt, og hvordan skal bearbeidelse og oppfølging av resultater organiseres?',
        'Er det behov for en ekstern konsulent, og i tilfelle i hvilke faser og for hvilke oppgaver?',
    ]:
        d.add_paragraph('☐  ' + t)
    small(d, 'De to listene over er gjengitt ordrett fra «Viktige momenter ved kartlegginger» i ' + QPS['source'].replace(' Vedlegg 2, s. 36–38.', ' Kapittel 7, s. 20.'))
    h(d, 'Det loven krever av kartleggingen', 2)
    for t in [
        'Kartleggingen og risikovurderingen gjøres i samarbeid med de ansatte og deres representanter (forskrift om utførelse av arbeid § 1A-2).',
        'Verneombudet er med i planleggingen og i vurderingen av resultatene.',
        'Risikovurderingen omfatter faktorene hver for seg, hvordan de påvirker hverandre, og den samlede risikoen, på kort og lang sikt (§ 1A-2).',
        'Det settes inn tiltak der risikoen tilsier det, og tiltakene følges opp (arbeidsmiljøloven § 3-1 andre ledd bokstav c).',
        'Kartleggingen gjentas regelmessig og ved endringer som kan ha betydning for risikoen, og alt dokumenteres (§ 1A-2).',
        'Anonymiteten er sikret: resultater vises bare for grupper som er store nok, og aldri for enkeltpersoner.',
    ]:
        d.add_paragraph('☐  ' + t)
    small(d, 'Kilder: arbeidsmiljøloven § 3-1 og forskrift om utførelse av arbeid kapittel 1A (Lovdata).')
    d.save(path)


FILES = {
    'orgpuls-risikovurdering-psykososialt-arbeidsmiljo.xlsx': risikovurdering,
    'orgpuls-tiltaksbank.xlsx': tiltaksbank,
    'orgpuls-arshjul-arbeidsmiljo.xlsx': arshjul,
    'qpsnordic-34-sporreskjema.docx': qps_docx,
    'qpsnordic-34-svarregistrering.xlsx': qps_xlsx,
    'orgpuls-informasjonsskriv-og-sjekkliste.docx': informasjonsskriv,
}

def printable(path: Path):
    """Every sheet prints one page wide, the wide ones in landscape."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        if ws.max_column > 4:
            ws.page_setup.orientation = 'landscape'
    wb.save(path)


if __name__ == '__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    for name, build in FILES.items():
        build(OUT / name)
        if name.endswith('.xlsx'):
            printable(OUT / name)
        print('wrote', OUT / name)
    with zipfile.ZipFile(OUT / 'orgpuls-kartleggingspakken.zip', 'w', zipfile.ZIP_DEFLATED) as z:
        for name in FILES:
            z.write(OUT / name, name)
    print('wrote', OUT / 'orgpuls-kartleggingspakken.zip')
